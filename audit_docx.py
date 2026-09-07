import zipfile, os, re, sys

path = r"C:\xampp\htdocs\ai_classroom_cheat_detection\AI_Classroom_Cheating_Detection_System_Complete_Documentation.docx"
print("File exists:", os.path.exists(path))
print("Size:", os.path.getsize(path))
try:
    z = zipfile.ZipFile(path)
    names = z.namelist()
    print("ZIP OK, entries:", len(names))
    for n in names:
        print(f"  {n} ({z.getinfo(n).file_size} bytes)")
    required = [
        "[Content_Types].xml",
        "_rels/.rels",
        "word/document.xml",
        "word/styles.xml",
        "word/settings.xml",
    ]
    for r in required:
        print(f"CHECK {r}:", "OK" if r in names else "MISSING")
    # check word/theme
    themes = [n for n in names if n.startswith("word/theme")]
    print("Themes:", themes)
    # check XML corruption for key files
    for f in required + themes:
        try:
            data = z.read(f)
            # try parse as XML
            import xml.etree.ElementTree as ET

            ET.fromstring(data)
            print(f"XML OK {f}")
        except Exception as e:
            print(f"XML FAIL {f}: {e}")
    # check document.xml for field codes
    doc_xml = z.read("word/document.xml").decode()
    print("\n--- Field codes search ---")
    if "fldChar" in doc_xml:
        print("Found fldChar - count", doc_xml.count("fldChar"))
        # extract snippet
        m = re.findall(r"fldChar.{0,80}", doc_xml)
        for x in m[:5]:
            print(x[:120])
    if "instrText" in doc_xml:
        print("Found instrText", doc_xml.count("instrText"))
    if "TOC" in doc_xml:
        print("Found TOC")
    # check invalid relationships
    try:
        rels = z.read("word/_rels/document.xml.rels").decode()
        print("\nRels entries", rels.count("<Relationship"))
    except Exception as e:
        print("No word/_rels/document.xml.rels", e)
    # try python-docx open
    try:
        from docx import Document

        d = Document(path)
        print(
            f"\npython-docx OPEN OK: paragraphs {len(d.paragraphs)}, tables {len(d.tables)}"
        )
    except Exception as e:
        print(f"python-docx FAIL: {e}")
        import traceback

        traceback.print_exc()
    # check for unsupported unicode
    try:
        txt = z.read("word/document.xml").decode()
        # find control chars
        bad = [c for c in txt if ord(c) < 32 and c not in "\n\r\t"]
        print("Unsupported control chars:", len(bad))
    except:
        pass
    # tables validation
    print("\n--- Table check ---")
    print("w:tbl count", doc_xml.count("w:tbl"))
finally:
    try:
        z.close()
    except:
        pass
