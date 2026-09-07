import zipfile, os, re

path = r"C:\xampp\htdocs\ai_classroom_cheat_detection\AI_Classroom_Cheating_Detection_System_Complete_Documentation_v2.docx"
print("Size", os.path.getsize(path))
z = zipfile.ZipFile(path)
names = z.namelist()
print("ZIP OK", len(names))
for req in [
    "[Content_Types].xml",
    "_rels/.rels",
    "word/document.xml",
    "word/styles.xml",
    "word/settings.xml",
    "word/theme/theme1.xml",
]:
    print(req, "OK" if req in names else "MISSING")
doc = z.read("word/document.xml").decode()
print("fldChar count", doc.count("fldChar"))
print(
    "invalid field hack present"
    if 'fldChar w:fldCharType="begin"/><w:instrText' in doc
    else "no raw fldChar hack"
)
print(
    "Unsupported control chars",
    len([c for c in doc if ord(c) < 32 and c not in "\n\r\t"]),
)
print("w:tbl", doc.count("w:tbl"))
try:
    import xml.etree.ElementTree as ET

    for n in names:
        if n.endswith(".xml"):
            ET.fromstring(z.read(n))
    print("XML all OK")
except Exception as e:
    print("XML FAIL", e)
from docx import Document

d = Document(path)
print("python-docx OPEN OK paragraphs", len(d.paragraphs), "tables", len(d.tables))
hs = [p.style.name for p in d.paragraphs if p.style.name.startswith("Heading")]
print("Headings", len(hs))
# verify not contain TOC field hack text
print("Contains TOC placeholder", "References → Table of Contents" in doc)
