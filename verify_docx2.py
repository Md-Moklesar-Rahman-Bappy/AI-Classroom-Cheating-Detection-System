from docx import Document

p = r"C:\xampp\htdocs\ai_classroom_cheat_detection\AI_Classroom_Cheating_Detection_System_Complete_Documentation.docx"
d = Document(p)
print("Paragraphs", len(d.paragraphs))
print("Tables", len(d.tables))
hs = [x.style.name for x in d.paragraphs if x.style.name.startswith("Heading")]
print("Headings", len(hs))
print(hs[:8])
# check secrets
txt = "\n".join([x.text for x in d.paragraphs])
secrets = ["APP_KEY", "password", "DB_PASSWORD", "camera", "RTSP"]
for s in secrets:
    if s.lower() in txt.lower() and "password `password` for local demo" in txt.lower():
        print("contains demo password mention (expected)")
# check Bengali
bengali_chars = sum(1 for c in txt if 0x0980 <= ord(c) <= 0x09FF)
print("Bengali chars", bengali_chars)
print("OK")
